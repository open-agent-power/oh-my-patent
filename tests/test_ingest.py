"""材料读取层（ingest）的验收测试。

这个文件要守住的不是「能转出文本」，而是**转换不会静默失真**。
理由：交给 Agent 的材料一旦缺了内容，它不会察觉——它会拿着少了公式、
塌了表格、丢了标题层级的文本去写权利要求，而且写得头头是道。

所以这里的断言分三类：

1. **公式不能丢位置**：``m:oMath`` 必须变成 ``⟪公式N: …⟫`` 占位，
   并且必须产生 ``FORMULA_APPROXIMATED`` 告警（占位是退化的，得让人知道）。
2. **表格必须是合法的 Markdown 表格**：有 ``| --- |`` 分隔行、数据行列对应。
3. **读不了的东西必须被点名**：.doc 报错、目录里的 .pdf/.xls 列出来，
   绝不能静默跳过。

第 2 类里有一条是回归守卫——``_promote_first_table_row``。mammoth 输出的
``<tr>`` 全是 ``<td>``，markdownify 会在前面补一个空表头行；没这一步修正，
每张表格头顶都会多出一行空白。
"""

from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from oh_my_patent import ingest


def _installed(module: str) -> bool:
    return importlib.util.find_spec(module) is not None


_HAS_MAMMOTH = _installed("mammoth") and _installed("markdownify")
_HAS_PPTX = _installed("pptx")
_HAS_PIL = _installed("PIL")

_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# --------------------------------------------------------------------------
# 造样本
# --------------------------------------------------------------------------


def _tiny_png(path: Path) -> Path:
    from PIL import Image

    Image.new("RGB", (12, 12), "white").save(path)
    return path


def _omath(xml_body: str) -> object:
    return parse_xml(f'<m:oMath {nsdecls("m")}>{xml_body}</m:oMath>')


def _omath_para(xml_body: str) -> object:
    return parse_xml(f'<m:oMathPara {nsdecls("m")}><m:oMath>{xml_body}</m:oMath></m:oMathPara>')


def _build_docx(path: Path, *, with_image: bool = True) -> Path:
    """造一份「像真实交底材料」的稿件：中文样式标题、表格、列表、公式、图片。"""
    doc = Document()

    # 中文样式名——中文 Word 用的就是「标题 1」，不是 "Heading 1"。
    doc.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("一种数据采集方法", style="标题 1")
    doc.add_paragraph("本发明涉及数据采集技术领域，具体而言涉及一种采集方法。")

    doc.add_paragraph("核心关系式如下：")
    # 块级公式 E = (1/2) m v
    holder = doc.add_paragraph()
    holder._p.append(
        _omath_para(
            "<m:r><m:t>E</m:t></m:r><m:r><m:t>=</m:t></m:r>"
            "<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>"
            "<m:r><m:t>m</m:t></m:r><m:r><m:t>v</m:t></m:r>"
        )
    )
    # 行内公式 m
    inline = doc.add_paragraph("其中质量 ")
    inline._p.append(_omath("<m:r><m:t>m</m:t></m:r>"))
    inline.add_run(" 的单位是千克。")

    doc.add_paragraph("各参数含义见下表：")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (left, right) in enumerate(
        [("参数", "含义"), ("m", "质量"), ("v", "速度")]
    ):
        table.cell(row, 0).text = left
        table.cell(row, 1).text = right

    doc.add_paragraph("第一步：采集原始数据", style="List Number")
    doc.add_paragraph("第二步：按上式换算", style="List Number")

    if with_image and _HAS_PIL:
        image_path = path.parent / "_tiny.png"
        _tiny_png(image_path)
        doc.add_picture(str(image_path))

    doc.save(path)
    return path


def _build_pptx(path: Path) -> Path:
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # 标题 + 内容

    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = "技术方案概述"
    body = first.placeholders[1].text_frame
    body.text = "第一，采集样本数据"
    body.add_paragraph().text = "第二，按公式计算"

    second = presentation.slides.add_slide(layout)
    second.shapes.title.text = "参数表"
    table = second.shapes.add_table(
        3, 2, Inches(1), Inches(2), Inches(6), Inches(2)
    ).table
    for row, pair in enumerate([("参数", "含义"), ("m", "质量"), ("v", "速度")]):
        for column, value in enumerate(pair):
            table.cell(row, column).text = value
    second.notes_slide.notes_text_frame.text = "强调单位统一"

    third = presentation.slides.add_slide(layout)
    third.shapes.title.text = "核心公式"
    frame = third.placeholders[1].text_frame
    frame.text = "能量满足 "
    frame.paragraphs[0]._p.append(
        etree.fromstring(
            f'<m:oMath xmlns:m="{_M}" xmlns:a="{_A}">'
            "<m:r><m:t>E</m:t></m:r>"
            "<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>"
            "</m:oMath>"
        )
    )

    presentation.slides.add_slide(layout)  # 纯空页：连标题都没有

    # 只有标题的一页。真实的评审 PPT 里，这类页面十有八九整页是一张架构图——
    # 图上没有可提取的文字，转出来只剩个标题。
    fifth = presentation.slides.add_slide(layout)
    fifth.shapes.title.text = "系统架构"

    presentation.save(path)
    return path


# --------------------------------------------------------------------------
# .docx
# --------------------------------------------------------------------------


@unittest.skipUnless(_HAS_MAMMOTH, "需要 mammoth 与 markdownify")
class TestDocxConversion(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls._tmp.name)
        cls.source = _build_docx(cls.root / "交底材料.docx")

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def _convert(self, **kwargs) -> ingest.IngestResult:
        media = self.root / f"media_{len(list(self.root.glob('media_*')))}"
        return ingest.docx_to_markdown(self.source, media_dir=media, **kwargs)

    # ---- 结构 ----

    def test_body_text_survives(self):
        result = self._convert()
        self.assertIn("本发明涉及数据采集技术领域", result.markdown)

    def test_chinese_style_name_becomes_heading(self):
        """中文样式名「标题 1」必须落成 Markdown 标题。

        mammoth 的默认 style map 只认英文样式名。实测：不加中文映射时，
        这一行会变成普通 ``<p>``——一份中文 Word 转出来会一个标题都不剩，
        全文层级塌成一堆并列段落。所以这条是硬要求，不是「锦上添花」。
        """
        result = self._convert()
        self.assertIn("# 一种数据采集方法", result.markdown)

    def test_list_survives(self):
        result = self._convert()
        self.assertIn("1. 第一步：采集原始数据", result.markdown)
        self.assertIn("2. 第二步：按上式换算", result.markdown)

    # ---- 公式：本模块存在的理由 ----

    def test_block_formula_becomes_placeholder(self):
        result = self._convert()
        self.assertIn("⟪公式1: E=12mv⟫", result.markdown)

    def test_inline_formula_keeps_its_slot(self):
        """行内公式最容易被无声吞掉：mammoth 原样输出会只剩一个空格。

        所以这里不满足于「有占位符」，而是要求它**长在句子中间**。
        """
        result = self._convert()
        self.assertIn("其中质量 ⟪公式2: m⟫ 的单位是千克。", result.markdown)

    def test_formula_triggers_a_warning(self):
        """占位文本是退化的（1/2 被摊成 12），必须告警，否则会被当原式照抄。"""
        result = self._convert()
        self.assertIn("FORMULA_APPROXIMATED", result.warning_codes())

    def test_formula_count_is_reported(self):
        result = self._convert()
        self.assertEqual(result.stats["formulas"], 2)

    def test_placeholder_avoids_markdown_link_syntax(self):
        """占位符刻意不用 ``[...]``——那在 Markdown 里是链接语法，会被解析歪。"""
        result = self._convert()
        self.assertNotIn("[公式1:", result.markdown)

    # ---- 表格 ----

    def test_table_is_a_real_markdown_table(self):
        result = self._convert()
        lines = result.markdown.split("\n")
        header = lines.index("| 参数 | 含义 |")
        self.assertEqual(lines[header + 1], "| --- | --- |")
        self.assertEqual(lines[header + 2], "| m | 质量 |")
        self.assertEqual(lines[header + 3], "| v | 速度 |")

    def test_no_empty_header_row(self):
        result = self._convert()
        for line in result.markdown.split("\n"):
            self.assertNotEqual(line.strip(), "|  |  |")

    # ---- 图片 ----

    @unittest.skipUnless(_HAS_PIL, "需要 Pillow 造样本图片")
    def test_images_are_exported_and_linked(self):
        result = self._convert()
        self.assertTrue(result.images)
        self.assertTrue(result.images[0].is_file())
        self.assertIn(result.images[0].name, result.markdown)

    @unittest.skipUnless(_HAS_PIL, "需要 Pillow 造样本图片")
    def test_skipping_images_is_announced(self):
        result = self._convert(extract_images=False)
        self.assertFalse(result.images)
        self.assertIn("IMAGE_SKIPPED", result.warning_codes())

    @unittest.skipUnless(_HAS_PIL, "需要 Pillow 造样本图片")
    def test_skipping_images_does_not_inline_base64(self):
        """``extract_images=False`` 曾经真的把图片转成 base64 塞进了正文。

        当时没有挂 image converter，mammoth 就退回默认行为：把每张图编码成
        ``data:image/png;base64,...`` 内联到 Markdown 里。一份带十几张截图的
        材料会膨胀成几 MB 的单行字符串，而参数名还写着「不导出图片」。
        这条断言守住「不导出就是真的不导出」。
        """
        result = self._convert(extract_images=False)
        self.assertNotIn("base64", result.markdown)
        self.assertNotIn("data:image", result.markdown)

    # ---- 头部注记 ----

    def test_header_carries_warnings_into_the_file(self):
        """告警必须写进 .md 本身——读文件的 Agent 看不到终端 stderr。"""
        result = self._convert()
        header = ingest.markdown_header(result)
        self.assertIn("FORMULA_APPROXIMATED", header)
        self.assertIn(self.source.name, header)


# --------------------------------------------------------------------------
# .pptx
# --------------------------------------------------------------------------


@unittest.skipUnless(_HAS_PPTX, "需要 python-pptx")
class TestPptxConversion(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls._tmp.name)
        cls.source = _build_pptx(cls.root / "评审材料.pptx")

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def _convert(self, **kwargs) -> ingest.IngestResult:
        return ingest.pptx_to_markdown(
            self.source, media_dir=self.root / "media", **kwargs
        )

    def test_slides_are_numbered(self):
        result = self._convert()
        for number in (1, 2, 3, 4):
            self.assertIn(f"## 第 {number} 页", result.markdown)

    def test_slide_text_survives(self):
        result = self._convert()
        self.assertIn("技术方案概述", result.markdown)
        self.assertIn("第一，采集样本数据", result.markdown)

    def test_table_has_separator_row(self):
        """少一行 ``| --- |``，渲染器就不认它是表格。

        上游同功能脚本恰好漏了这一行（``pptx_to_md.py``），输出的是
        「看起来像表格、实际渲染成一行竖线」的东西。
        """
        result = self._convert()
        lines = result.markdown.split("\n")
        header = lines.index("| 参数 | 含义 |")
        self.assertEqual(lines[header + 1], "| --- | --- |")
        self.assertIn("| m | 质量 |", lines)

    def test_notes_are_kept(self):
        result = self._convert()
        self.assertIn("强调单位统一", result.markdown)

    def test_inline_formula_becomes_placeholder(self):
        result = self._convert()
        self.assertIn("⟪公式1: E12⟫", result.markdown)
        self.assertIn("FORMULA_APPROXIMATED", result.warning_codes())

    def test_textless_slide_is_called_out(self):
        """纯图形页转出来是空的，必须点名——否则 Agent 会以为这页没内容。"""
        result = self._convert()
        self.assertIn("SLIDE_EMPTY", {n.code for n in result.notices})
        self.assertEqual(result.stats["slides_without_text"], 1)

    def test_title_only_slide_is_called_out(self):
        """只有标题、没有正文的一页也要点名。

        最初只判「一个字都没有」的页面，结果「系统架构」这种页面被漏掉了：
        它有标题，所以不算空；但正文全在一张图上，图是提取不出来的。
        验收时就是这一页让问题暴露出来的。
        """
        result = self._convert()
        self.assertIn("SLIDE_TITLE_ONLY", {n.code for n in result.notices})
        self.assertEqual(result.stats["slides_title_only"], 1)

    def test_pptx_header_lists_both_warnings_and_infos(self):
        """端到端：真实夹具转出的 .md，两类提示都必须出现在 header 里。

        上面两条只断言了 notice 被**产生**。产生不等于**送达**——
        Agent 手里只有这个 .md 文件，它看不到终端的提示输出。
        """
        header = ingest.markdown_header(self._convert())
        self.assertIn("FORMULA_APPROXIMATED", header)
        self.assertIn("SLIDE_TITLE_ONLY", header)
        self.assertIn("SLIDE_EMPTY", header)

    def test_pptx_file_really_carries_the_notice(self):
        """整条链路：转换 → 拼 header → Agent 实际读到的文本。"""
        result = self._convert()
        text = ingest.markdown_header(result) + result.markdown
        self.assertIn("SLIDE_TITLE_ONLY", text)
        self.assertIn("系统架构", text)


# --------------------------------------------------------------------------
# 错误路径与分派
# --------------------------------------------------------------------------


class TestErrorsAndDispatch(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_missing_file_is_reported(self):
        with self.assertRaises(ingest.IngestError):
            ingest.docx_to_markdown(self.root / "不存在.docx")

    def test_legacy_doc_is_refused_with_advice(self):
        legacy = self.root / "旧稿.doc"
        legacy.write_bytes(b"not really a doc")
        with self.assertRaises(ingest.IngestError) as ctx:
            ingest.convert(legacy)
        self.assertIn("另存为", str(ctx.exception))

    def test_legacy_ppt_is_refused(self):
        legacy = self.root / "旧演示.ppt"
        legacy.write_bytes(b"x")
        with self.assertRaises(ingest.IngestError):
            ingest.convert(legacy)

    def test_unknown_suffix_is_refused(self):
        other = self.root / "说明.txt"
        other.write_text("hi", encoding="utf-8")
        with self.assertRaises(ingest.IngestError):
            ingest.convert(other)

    def test_wrong_suffix_for_docx_reader(self):
        other = self.root / "表格.xlsx"
        other.write_bytes(b"x")
        with self.assertRaises(ingest.IngestError):
            ingest.docx_to_markdown(other)

    @unittest.skipUnless(_HAS_MAMMOTH, "需要 mammoth 与 markdownify")
    def test_convert_dispatches_by_suffix(self):
        source = _build_docx(self.root / "分派.docx", with_image=False)
        result = ingest.convert(source, media_dir=self.root / "m")
        self.assertIn("本发明涉及数据采集技术领域", result.markdown)


# --------------------------------------------------------------------------
# 纯函数
# --------------------------------------------------------------------------


class TestTableHeaderPromotion(unittest.TestCase):
    def test_empty_header_is_replaced_by_first_row(self):
        raw = "|  |  |\n| --- | --- |\n| 参数 | 含义 |\n| m | 质量 |"
        got = ingest._promote_first_table_row(raw)
        self.assertEqual(
            got.split("\n"),
            ["| 参数 | 含义 |", "| --- | --- |", "| m | 质量 |"],
        )

    def test_real_header_is_left_alone(self):
        raw = "| 参数 | 含义 |\n| --- | --- |\n| m | 质量 |"
        self.assertEqual(ingest._promote_first_table_row(raw), raw)

    def test_single_column_table(self):
        raw = "|  |\n| --- |\n| 甲 |"
        self.assertEqual(ingest._promote_first_table_row(raw).split("\n")[0], "| 甲 |")


class TestIngestBookkeeping(unittest.TestCase):
    def test_warning_codes_are_deduplicated_and_filtered(self):
        result = ingest.IngestResult(
            source=Path("x.docx"),
            markdown="",
            notices=[
                ingest.Notice("info", "SLIDE_EMPTY", "空"),
                ingest.Notice("warning", "FORMULA_APPROXIMATED", "公式"),
            ],
        )
        self.assertEqual(result.warning_codes(), {"FORMULA_APPROXIMATED"})
        self.assertEqual(len(result.warnings), 1)

    def test_header_without_notices_stays_quiet(self):
        result = ingest.IngestResult(
            source=Path("x.docx"), markdown="", stats={"formulas": 0}
        )
        header = ingest.markdown_header(result)
        self.assertNotIn("转换告警", header)
        self.assertNotIn("其它提示", header)

    def test_info_notice_alone_still_reaches_the_header(self):
        """没有 warning、只有 info 时，header 也必须把 info 说出来。

        回归守卫。header 原先是 ``if result.has_warnings:`` 才写提示——
        于是「一份文件没有任何 warning」就直接什么都不写，而
        ``SLIDE_TITLE_ONLY`` / ``SLIDE_EMPTY`` 说的恰恰是
        「**有内容我没能给你**」。实测一份 PPT 的第 4 页只剩「系统架构」
        四个字，Agent 读到的是一页空内容，根本不会想到它原本是整张架构图。
        对只有 Read 权限的 Agent 来说，「我没读到」和「材料里没有」
        必须能被区分开。
        """
        result = ingest.IngestResult(
            source=Path("x.docx"),
            markdown="",
            notices=[ingest.Notice("info", "SLIDE_TITLE_ONLY", "只有标题")],
        )
        header = ingest.markdown_header(result)
        self.assertIn("SLIDE_TITLE_ONLY", header)
        self.assertNotIn("转换告警", header)

    def test_supported_suffixes_cover_the_office_formats(self):
        for suffix in (".docx", ".pptx", ".ppsx", ".docm"):
            self.assertIn(suffix, ingest.SUPPORTED_SUFFIXES)


if __name__ == "__main__":
    unittest.main()
