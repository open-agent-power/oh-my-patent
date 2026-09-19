"""渲染层的验收测试。

这个文件的核心不是「能生成文件」，而是**生成的文件在 XML 层面是合法的**。
理由：WPS 对 OOXML 相当宽容，Office 却会严格按 schema 校验——顺序错、
属性缺失都可能让 Office 弹「文件已损坏，是否修复」。只在 WPS 里点开看一眼，
根本发现不了这类问题。所以这里直接把 .docx 当 zip 拆开逐项检查。
"""

from __future__ import annotations

import posixpath
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.shared import Mm

from oh_my_patent.oxml import _PPR_ORDER, _RPR_ORDER, _qualified
from oh_my_patent.parser import parse_markdown
from oh_my_patent.render import DocxRenderer, render_docx
from oh_my_patent.spec import Font, RenderOptions

REPO_ROOT = Path(__file__).resolve().parent.parent
EXAMPLE = REPO_ROOT / "examples" / "invention-example.md"

#: 命名空间，用于 ElementTree 查询。
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: A4 与规定的页边距，换算成 twips（1 毫米 = 56.6929 twips）。
EXPECTED_PG_SZ = (11906, 16838)          # 210mm x 297mm
EXPECTED_PG_MAR = {                        # 上25 右15 下15 左25 毫米
    "top": 1417,
    "right": 850,
    "bottom": 850,
    "left": 1417,
}


def _minimal_draft():
    return parse_markdown(
        "---\ntitle: 一种测试用装置\ntype: invention\n---\n\n"
        "# 权利要求书\n\n"
        "1. 一种测试用装置，其特征在于，包括壳体。\n"
        "2. 根据权利要求1所述的测试用装置，其特征在于，所述壳体为金属壳。\n\n"
        "# 说明书\n\n"
        "## 技术领域\n本发明涉及测试装置技术领域。\n\n"
        "## 背景技术\n现有装置存在缺陷。\n\n"
        "## 发明内容\n本发明提供一种测试用装置。\n\n"
        "## 具体实施方式\n本实施例中，壳体采用铝合金制成。\n\n"
        "# 摘要\n本发明公开一种测试用装置，结构简单。\n"
    )


class DocxStructureTest(unittest.TestCase):
    """把生成的 .docx 当 zip 拆开，逐项检查结构合法性。"""

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls._tmp.name) / "out.docx"
        render_docx(_minimal_draft(), cls.path)
        cls.zip = zipfile.ZipFile(cls.path)
        cls.parts = {
            name: cls.zip.read(name)
            for name in cls.zip.namelist()
            if name.endswith(".xml")
        }

    @classmethod
    def tearDownClass(cls):
        cls.zip.close()
        cls._tmp.cleanup()

    # ---- 基本完整性 ----

    def test_ooxml_parts_present(self):
        """OOXML 必需的部分一个都不能少。"""
        required = {
            "[Content_Types].xml",
            "_rels/.rels",
            "word/document.xml",
            "word/styles.xml",
            "word/settings.xml",
            "word/_rels/document.xml.rels",
        }
        missing = required - set(self.zip.namelist())
        self.assertFalse(missing, f"缺少 OOXML 必需部分：{missing}")

    def test_all_xml_parts_well_formed(self):
        """每个 XML 部分都能被解析，说明没有坏字节。"""
        for name, payload in self.parts.items():
            with self.subTest(part=name):
                ElementTree.fromstring(payload)

    def test_content_types_covers_every_part(self):
        """[Content_Types].xml 必须为每个 part 声明类型。

        漏声明时 WPS 往往照样打开，Office 则可能判定文件损坏。
        """
        types = ElementTree.fromstring(self.parts["[Content_Types].xml"])
        default_exts = {
            el.get("Extension").lower()
            for el in types.findall(
                "{http://schemas.openxmlformats.org/package/2006/content-types}Default"
            )
        }
        overrides = {
            el.get("PartName")
            for el in types.findall(
                "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
            )
        }
        for name in self.zip.namelist():
            if name == "[Content_Types].xml" or name.endswith("/"):
                continue
            with self.subTest(part=name):
                if "/" + name in overrides:
                    continue
                ext = name.rsplit(".", 1)[-1].lower()
                self.assertIn(ext, default_exts, f"{name} 既无 Override 也无 Default 声明")

    def test_relationships_resolve(self):
        """每个关系的 Target 都要真实存在。"""
        rels = ElementTree.fromstring(self.zip.read("word/_rels/document.xml.rels"))
        names = set(self.zip.namelist())
        for rel in rels:
            target = rel.get("Target")
            if rel.get("TargetMode") == "External" or target.startswith("http"):
                continue
            # 关系里的 Target 允许写相对路径（如 ../customXml/item1.xml），
            # 必须先用 posixpath 归一化再比对包内条目名。
            resolved = posixpath.normpath(posixpath.join("word", target))
            with self.subTest(target=target):
                self.assertIn(resolved, names, f"关系指向了不存在的部分：{target} → {resolved}")

    # ---- WPS / Office 双适配的关键断言 ----

    def _styles(self):
        return ElementTree.fromstring(self.parts["word/styles.xml"])

    def test_no_theme_references_anywhere(self):
        """全文档不得残留主题字体或主题色引用。

        这是 WPS 与 Office 版面产生差异的根因：主题里
        ``<a:ea typeface=""/>`` 是空的，中文字体只写在按脚本区分的
        ``<a:font script="Hans"/>`` 中，解析结果由各软件自行决定。
        """
        pattern = re.compile(
            rb"w:(asciiTheme|hAnsiTheme|eastAsiaTheme|cstheme|themeColor|themeFill|themeTint|themeShade)"
        )
        for name, payload in self.parts.items():
            with self.subTest(part=name):
                found = set(pattern.findall(payload))
                self.assertFalse(found, f"{name} 残留主题引用：{found}")

    def test_doc_defaults_declare_east_asian_font(self):
        """docDefaults 必须显式声明 ``w:eastAsia``。

        缺了这一项，Office 会用主题的 minorEastAsia（常解析为等线），
        WPS 会用宋体，同一份文件两边排版不同。
        """
        rpr_default = self._styles().find(f".//{W}rPrDefault/{W}rPr")
        self.assertIsNotNone(rpr_default)
        rfonts = rpr_default.find(f"{W}rFonts")
        self.assertIsNotNone(rfonts, "docDefaults 里没有 w:rFonts")
        self.assertEqual(rfonts.get(f"{W}eastAsia"), Font.SONG.value)
        self.assertEqual(rfonts.get(f"{W}hint"), "eastAsia")
        for attr in (f"{W}ascii", f"{W}hAnsi", f"{W}cs"):
            self.assertTrue(rfonts.get(attr), f"w:rFonts 缺少 {attr}")

    def test_every_run_declares_both_fonts(self):
        """每个 run 都要同时写死中文字体与西文字体。"""
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])
        runs = doc_xml.findall(f".//{W}r")
        self.assertTrue(runs)
        for run in runs:
            rfonts = run.find(f"{W}rPr/{W}rFonts")
            with self.subTest(text="".join(t.text or "" for t in run.findall(f"{W}t"))[:20]):
                self.assertIsNotNone(rfonts, "run 没有显式字体声明")
                self.assertEqual(rfonts.get(f"{W}eastAsia"), Font.SONG.value)
                self.assertTrue(rfonts.get(f"{W}ascii"))

    def test_run_properties_follow_schema_order(self):
        """``w:rPr`` 的子元素顺序必须符合 ECMA-376。

        顺序错了 Office 可能直接报文件损坏。这个断言专门守住
        `insert_ordered` 不被改坏。
        """
        order = [_qualified(tag) for tag in _RPR_ORDER]
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])

        checked = 0
        for rpr in doc_xml.iter(f"{W}rPr"):
            positions = []
            for child in rpr:
                if child.tag in order:
                    positions.append(order.index(child.tag))
            with self.subTest(children=[c.tag.split("}")[-1] for c in rpr]):
                self.assertEqual(positions, sorted(positions), "w:rPr 子元素顺序不符合 schema")
            checked += 1
        self.assertGreater(checked, 0)

    def test_paragraph_properties_follow_schema_order(self):
        """``w:pPr`` 的子元素顺序同样必须符合 schema。"""
        order = [_qualified(tag) for tag in _PPR_ORDER]
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])

        checked = 0
        for ppr in doc_xml.iter(f"{W}pPr"):
            positions = []
            for child in ppr:
                if child.tag in order:
                    positions.append(order.index(child.tag))
            with self.subTest(children=[c.tag.split("}")[-1] for c in ppr]):
                self.assertEqual(positions, sorted(positions), "w:pPr 子元素顺序不符合 schema")
            checked += 1
        self.assertGreater(checked, 0)

    def test_page_setup_matches_regulation(self):
        """纸张 A4、页边距上25 右15 下15 左25 毫米。"""
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])
        sect_pr = doc_xml.find(f".//{W}sectPr")
        self.assertIsNotNone(sect_pr)

        pg_sz = sect_pr.find(f"{W}pgSz")
        self.assertEqual(
            (int(pg_sz.get(f"{W}w")), int(pg_sz.get(f"{W}h"))), EXPECTED_PG_SZ
        )

        pg_mar = sect_pr.find(f"{W}pgMar")
        for side, expected in EXPECTED_PG_MAR.items():
            with self.subTest(side=side):
                self.assertEqual(int(pg_mar.get(f"{W}{side}")), expected)

    def test_document_grid_disabled(self):
        """必须关闭文档网格。

        启用行网格时行距会被吸附到网格上，而 WPS 与 Office 的吸附算法
        不同，会导致两边每页行数不一致。关掉后行距只由段落自身决定。
        """
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])
        doc_grid = doc_xml.find(f".//{W}docGrid")
        self.assertIsNotNone(doc_grid)
        self.assertEqual(doc_grid.get(f"{W}type"), "default")
        self.assertIsNone(doc_grid.get(f"{W}linePitch"))

    def test_compatibility_mode_is_declared(self):
        """settings.xml 必须声明 compatibilityMode=15。

        缺失时 WPS 可能按 Word 2003 兼容模式排版，行距与字距都会变。
        """
        settings = ElementTree.fromstring(self.parts["word/settings.xml"])
        modes = [
            el.get(f"{W}val")
            for el in settings.iter(f"{W}compatSetting")
            if el.get(f"{W}name") == "compatibilityMode"
        ]
        self.assertEqual(modes, ["15"])

    def test_first_line_indent_uses_character_unit(self):
        """正文首行缩进必须用「字符」单位（``w:firstLineChars``）。

        用磅值缩进的话，字号一变缩进就不对了；中文文书的缩进量要求随字号变化。
        """
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])
        indents = list(doc_xml.iter(f"{W}ind"))
        self.assertTrue(indents)
        char_based = [el for el in indents if el.get(f"{W}firstLineChars") == "200"]
        self.assertTrue(char_based, "没有任何段落使用 2 字符的首行缩进")

    def test_no_automatic_numbering(self):
        """不得使用自动编号列表。

        WPS 对 numbering.xml 的实现与 Office 有差异，会用自动编号把
        权利要求编号渲染成圆点或错位。这里全部用纯文本编号，所以
        document.xml 里不应出现 ``w:numPr``。
        """
        doc_xml = ElementTree.fromstring(self.parts["word/document.xml"])
        self.assertEqual(list(doc_xml.iter(f"{W}numPr")), [])


class DocxContentTest(unittest.TestCase):
    """检查内容与排版语义是否正确。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.out = Path(self._tmp.name)

    def _render(self, draft, **kwargs):
        options = kwargs.pop("options", None)
        path = self.out / "x.docx"
        render_docx(draft, path, options=options)
        return Document(str(path))

    def test_text_content_survives_roundtrip(self):
        """生成的文字必须能原样读回来。"""
        document = self._render(_minimal_draft())
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("一种测试用装置", text)
        self.assertIn("1. 一种测试用装置，其特征在于，包括壳体。", text)
        self.assertIn("【技术领域】", text)
        self.assertIn("【具体实施方式】", text)

    def test_claims_come_first_then_description_then_abstract(self):
        """各部分顺序符合单行本的阅读顺序。"""
        document = self._render(_minimal_draft())
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertLess(text.index("权利要求书"), text.index("【技术领域】"))
        self.assertLess(text.index("【技术领域】"), text.index("说明书摘要"))

    def test_page_break_between_parts(self):
        """各部分之间用 ``w:pageBreakBefore`` 分页，而不是插入空段落。

        插入空段落会在上一部分末尾留下空行，在 WPS 里有时被撑成整整一页。
        """
        path = self.out / "pb.docx"
        render_docx(_minimal_draft(), path)
        document = Document(str(path))
        breaks = [
            p.text[:12]
            for p in document.paragraphs
            if p._element.find(f"{W}pPr/{W}pageBreakBefore") is not None
        ]
        self.assertEqual(len(breaks), 2, f"分页位置不对：{breaks}")
        self.assertIn("说明书摘要", breaks[1])

    def test_no_page_breaks_when_disabled(self):
        options = RenderOptions()
        options.page_break_between_parts = False
        path = self.out / "nopb.docx"
        render_docx(_minimal_draft(), path, options=options)
        document = Document(str(path))
        breaks = [
            p for p in document.paragraphs
            if p._element.find(f"{W}pPr/{W}pageBreakBefore") is not None
        ]
        self.assertEqual(breaks, [])

    def test_font_option_takes_effect(self):
        """换成黑体后，字体声明要跟着变。"""
        options = RenderOptions()
        options.font = Font.HEI
        path = self.out / "hei.docx"
        render_docx(_minimal_draft(), path, options=options)

        document = Document(str(path))
        rfonts = document.paragraphs[0].runs[0]._element.find(f"{W}rPr/{W}rFonts")
        self.assertEqual(rfonts.get(f"{W}eastAsia"), Font.HEI.value)
        # 黑体配的西文字体是 Arial
        self.assertEqual(rfonts.get(f"{W}ascii"), Font.HEI.western)

    def test_split_produces_separate_documents(self):
        """拆分模式下每个部分独立成文件。

        这份测试稿没有附图，所以「说明书附图」部分没有内容、不会产出文件，
        最终是权利要求书、说明书、说明书摘要三份。
        """
        renderer = DocxRenderer(_minimal_draft())
        written = renderer.save_split(self.out)
        names = sorted(p.name for p in written)
        self.assertEqual(len(written), 3, names)
        self.assertTrue(any("权利要求书" in n for n in names))
        self.assertTrue(any("说明书摘要" in n for n in names))
        for path in written:
            document = Document(str(path))
            self.assertTrue(document.paragraphs)

    def test_trailing_empty_paragraph_removed(self):
        """文末不留空段落。"""
        document = self._render(_minimal_draft())
        last = document.paragraphs[-1]
        self.assertTrue(last.text.strip(), "文末残留空段落")


class SubsectionStyleTest(unittest.TestCase):
    """「发明内容」三段式在不同处理方式下的输出。"""

    SOURCE = (
        "---\ntitle: 一种示例方法\ntype: invention\n---\n\n"
        "# 说明书\n\n"
        "## 发明内容\n\n"
        "### 要解决的技术问题\n现有方案成本高。\n\n"
        "### 技术方案\n采用替代材料。\n\n"
        "### 有益效果\n成本下降三成。\n"
    )

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.draft = parse_markdown(self.SOURCE)

    def _text(self, style):
        options = RenderOptions()
        options.subsection_style = style
        path = Path(self._tmp.name) / f"{style.value}.docx"
        render_docx(self.draft, path, options=options)
        return "\n".join(p.text for p in Document(str(path)).paragraphs)

    def test_inline_prepends_bold_lead_in(self):
        from oh_my_patent.spec import SubsectionStyle

        text = self._text(SubsectionStyle.INLINE)
        self.assertIn("技术方案：采用替代材料。", text)
        self.assertNotIn("技术方案\n", text)

    def test_keep_preserves_headings(self):
        from oh_my_patent.spec import SubsectionStyle

        text = self._text(SubsectionStyle.KEEP)
        self.assertIn("技术方案", text.splitlines())
        self.assertIn("有益效果", text.splitlines())

    def test_flatten_drops_headings(self):
        from oh_my_patent.spec import SubsectionStyle

        text = self._text(SubsectionStyle.FLATTEN)
        self.assertNotIn("技术方案：", text)
        self.assertNotIn("有益效果", text)
        self.assertIn("采用替代材料。", text)


class ImageRenderingTest(unittest.TestCase):
    """附图插入。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.out = Path(self._tmp.name)

    def _make_png(self, name: str, width: int, height: int) -> Path:
        """用 zlib 手工拼一张 PNG，避免引入 Pillow 依赖。"""
        import struct
        import zlib

        raw = b"".join(
            b"\x00" + bytes([200, 30, 30] * width) for _ in range(height)
        )

        def chunk(tag: bytes, payload: bytes) -> bytes:
            body = tag + payload
            return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

        png = (
            b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b"")
        )
        path = self.out / name
        path.write_bytes(png)
        return path

    def test_image_appears_in_abstract_and_drawings(self):
        """有附图时，同一幅图会同时出现在摘要与说明书附图两部分。

        这是符合《专利法实施细则》的做法：有附图的申请必须提供一幅最能
        说明技术特征的附图作为摘要附图，而它本身仍是说明书附图之一。
        """
        png = self._make_png("fig1.png", 200, 400)  # 宽高比 1:2，必然触及高度上限
        draft = parse_markdown(
            "---\ntitle: 一种示例装置\ntype: invention\n---\n\n"
            "# 说明书\n\n## 技术领域\n本发明涉及示例装置。\n\n"
            "## 具体实施方式\n如图1所示。\n\n"
            "# 摘要\n本发明公开一种示例装置。\n\n"
            "# 附图\n\n"
            f"![图1]({png.as_posix()})\n"
        )
        path = self.out / "withimg.docx"
        render_docx(draft, path)

        document = Document(str(path))
        shapes = document.inline_shapes
        self.assertEqual(len(shapes), 2, "应分别落在摘要与说明书附图")

        for shape in shapes:
            with self.subTest(height=shape.height):
                self.assertLessEqual(shape.height, Mm(200.1))
                self.assertGreater(shape.height, Mm(1))

        texts = [p.text for p in document.paragraphs]
        self.assertLess(texts.index("说明书摘要"), texts.index("说明书附图"))

    def test_empty_part_is_not_written_in_split_mode(self):
        """没有内容的文书部分不应产出空文件。

        稿件注册了附图但没有实际图片时，「说明书附图」这一部分没有任何
        可写的内容，产出空文件只会让提交环节多一份无用材料。
        """
        draft = parse_markdown(
            "---\ntitle: 一种示例装置\ntype: invention\n---\n\n"
            "# 权利要求书\n\n1. 一种示例装置，其特征在于，包括壳体。\n\n"
            "# 说明书\n\n## 技术领域\n本发明涉及示例装置。\n\n"
            "## 具体实施方式\n如图1所示。\n"
        )
        renderer = DocxRenderer(draft)
        written = renderer.save_split(self.out)
        names = [p.name for p in written]
        self.assertFalse(any("说明书附图" in n for n in names), names)

    def test_missing_image_does_not_crash(self):
        """图片文件缺失时应降级为提示文字，而不是抛异常中断生成。"""
        draft = parse_markdown(
            "---\ntitle: 一种示例装置\ntype: invention\n---\n\n"
            "# 说明书\n\n## 技术领域\n本发明涉及示例装置。\n\n"
            "# 附图\n\n![图1](不存在的文件.png)\n"
        )
        path = self.out / "broken.docx"
        written = render_docx(draft, path)
        self.assertTrue(written[0].exists())
        text = "\n".join(p.text for p in Document(str(path)).paragraphs)
        self.assertIn("附图缺失", text)


if __name__ == "__main__":
    unittest.main()
